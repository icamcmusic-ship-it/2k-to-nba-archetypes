#!/usr/bin/env python3
"""Extract the NBA 2K archetype database from the 2K_to_BBGM_Archetypes.xlsx
workbook into a JS data file consumed by the web app (docs/archetypes.js).

Usage: python3 tools/extract_archetypes.py path/to/2K_to_BBGM_Archetypes.xlsx
"""
import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

# 2K stat columns used for matching, in output order. Columns are located by
# header text rather than by index: hardcoded indices silently mis-map every
# stat if a column is ever inserted into the workbook.
STAT_NAMES = [
    "Close Shot", "Driving Layup", "Driving Dunk", "Standing Dunk", "Post Control",
    "Midrange Shot", "Threepoint Shot", "Free Throw", "Pass Accuracy", "Ball Handle",
    "Speed With Ball", "Interior Defense", "Perimeter Defense", "Steal", "Block",
    "Offensive Rebound", "Defensive Rebound", "Speed", "Agility", "Strength", "Vertical",
]

# header text (normalized) -> canonical stat name
HEADER_ALIASES = {
    "close shot": "Close Shot", "driving layup": "Driving Layup",
    "driving dunk": "Driving Dunk", "standing dunk": "Standing Dunk",
    "post control": "Post Control", "midrange shot": "Midrange Shot",
    "mid-range shot": "Midrange Shot", "midrange": "Midrange Shot",
    "threepoint shot": "Threepoint Shot", "three point shot": "Threepoint Shot",
    "three-point shot": "Threepoint Shot", "3pt": "Threepoint Shot",
    "free throw": "Free Throw", "pass accuracy": "Pass Accuracy",
    "ball handle": "Ball Handle", "speed with ball": "Speed With Ball",
    "interior defense": "Interior Defense", "interior d": "Interior Defense",
    "perimeter defense": "Perimeter Defense", "perimeter d": "Perimeter Defense",
    "steal": "Steal", "block": "Block",
    "offensive rebound": "Offensive Rebound", "defensive rebound": "Defensive Rebound",
    "speed": "Speed", "agility": "Agility", "strength": "Strength", "vertical": "Vertical",
}


def norm(v):
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def find_columns(ws):
    """Locate each stat column by its header text, scanning the first few rows."""
    for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
        found = {}
        for idx, cell in enumerate(row):
            stat = HEADER_ALIASES.get(norm(cell))
            if stat and stat not in found:
                found[stat] = idx
        if len(found) == len(STAT_NAMES):
            return found
    missing = [s for s in STAT_NAMES if s not in found]
    raise SystemExit(
        "Could not locate all stat columns by header text in the Archetype_DB "
        "sheet. Missing: " + ", ".join(missing))

POS_ABBR = {
    "Point Guard": "PG",
    "Shooting Guard": "SG",
    "Small Forward": "SF",
    "Power Forward": "PF",
    "Center": "C",
}

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = REPO_ROOT / "docs" / "archetypes.js"


def parse_range(v):
    """Return (min, max) of a '70-85' range, or (v, v) for a bare number."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return (float(v), float(v))
    s = str(v).strip()
    m = re.match(r"^(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)$", s)
    if m:
        lo, hi = float(m.group(1)), float(m.group(2))
        return (lo, hi) if lo <= hi else (hi, lo)
    try:
        n = float(s)
        return (n, n)
    except ValueError:
        return None


def parse_height(v):
    """Parse '6\'4" to 6\'6"' or '6\'6"' into (min_inches, max_inches)."""
    if v is None:
        return None, None
    feet_in = re.findall(r"(\d+)'\s*(\d+)?\"?", str(v))
    if not feet_in:
        return None, None
    inches = [int(f) * 12 + int(i or 0) for f, i in feet_in]
    return min(inches), max(inches)


def main():
    parser = argparse.ArgumentParser(
        description="Extract the NBA 2K archetype database into docs/archetypes.js")
    parser.add_argument("workbook", help="path to 2K_to_BBGM_Archetypes.xlsx")
    parser.add_argument("-o", "--output", default=str(DEFAULT_OUTPUT),
                         help=f"output JS path (default: {DEFAULT_OUTPUT})")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    ws = wb["Archetype_DB"]
    stat_cols = find_columns(ws)

    archetypes = []
    skipped = 0
    for row_num, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        name = row[0]
        if not name:
            continue
        hmin, hmax = parse_height(row[2])
        stats = {}
        ok = True
        bad_stat = None
        for stat, col in stat_cols.items():
            val = parse_range(row[col])
            if val is None:
                ok = False
                bad_stat = stat
                break
            stats[stat] = (round(val[0], 2), round(val[1], 2))
        if not ok or hmin is None:
            skipped += 1
            reason = f"unparseable '{bad_stat}'" if not ok else "unparseable height"
            print(f"  skipping row {row_num} ({name!r}): {reason}", file=sys.stderr)
            continue
        archetypes.append({
            "name": name,
            "pos": POS_ABBR.get(row[1], row[1]),
            "hMin": hmin,
            "hMax": hmax,
            # each stat is [min, max] of the build's legal range
            "stats": [list(stats[s]) for s in STAT_NAMES],
        })

    out = {
        "statNames": STAT_NAMES,
        "archetypes": archetypes,
    }
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w") as f:
        f.write("// Generated by tools/extract_archetypes.py — do not edit by hand\n")
        f.write("const ARCHETYPE_DATA = ")
        json.dump(out, f, separators=(",", ":"))
        f.write(";\n")
        # also require()-able from Node so the test suite and CI can load the
        # same data file the browser does
        f.write('if (typeof module !== "undefined" && module.exports) module.exports = ARCHETYPE_DATA;\n')
    print(f"Wrote {len(archetypes)} archetypes to {output_path}"
          + (f" ({skipped} row(s) skipped, see above)" if skipped else ""))
    if skipped:
        # a silently-dropped build is a data-quality regression, not a warning
        print(f"error: {skipped} row(s) could not be parsed (listed above)", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main() or 0)
